#!/usr/bin/env python3
"""
24_build_ndt_docx.py

Builds the Nephrology Dialysis Transplantation submission files.

Requirements taken from the live Author Guidelines and Information to Authors
(academic.oup.com/ndt, read 21 September 2026):
  - main text maximum 3,500 words INCLUDING a 300-word abstract, excluding references,
    tables and figures
  - structured abstract with the four headings Background and hypothesis / Methods /
    Results / Conclusions
  - mandatory Key Learning Points: What was known / This study adds / Potential impact,
    at most 3 bullets and 50 words each
  - at most 5 keywords, at most 50 references, running head at most 50 characters
  - section order: Introduction, Materials and Methods, Results, Discussion,
    Acknowledgements, Conflict of Interest Statement, Authors' Contributions, Funding,
    Data Availability Statement, References, Tables, Figure Legends, Figures
  - numbered references in order of appearance, Vancouver style, cited as superscripts
  - headings on separate lines and not numbered; all pages numbered consecutively

Citation numbers are produced by 16_references.py from the [PMID ...] and [REF:...] tags,
never typed, and are rendered here as superscripts.

Three files are written:
  NDT_submission_manuscript.docx   text and tables, no embedded figures
  NDT_review_copy.docx             the same with figures embedded, for reading
  NDT_supplementary_methods.docx   the supplementary methods

Author: Christopher Lawrence
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
PUB = ROOT / "publication"
FIG = PUB / "figures_ejhg"
SRC = PUB / "NDT_manuscript_v1.0.referenced.md"
SUPP = PUB / "NDT_supplementary_methods.md"
TABLES = PUB / "tables_ejhg.md"
INK = RGBColor(0, 0, 0)
FIGURE_FILES = ["figure1_discrimination.png", "figure2_avi_attribution.png", "figure3_locus_rank.png"]

# NDT's prescribed order. The abstract and key learning points precede the Introduction.
ORDER = ["Abstract", "Key learning points", "Introduction", "Materials and Methods", "Results",
         "Discussion", "Acknowledgements", "Conflict of Interest Statement",
         "Authors' Contributions", "Funding", "Data Availability Statement", "References",
         "Figure Legends"]


def add_line_numbers(section) -> None:
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    section._sectPr.append(ln)


def add_page_numbers(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for instr, _ in (("begin", None), (None, "PAGE"), ("end", None)):
        el = OxmlElement("w:fldChar") if instr else OxmlElement("w:instrText")
        if instr:
            el.set(qn("w:fldCharType"), instr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        run._r.append(el)


def style(doc: Document) -> None:
    n = doc.styles["Normal"]
    n.font.name, n.font.size, n.font.color.rgb = "Times New Roman", Pt(12), INK
    pf = n.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(0)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.25)
        s.top_margin = s.bottom_margin = Inches(1.0)
        add_line_numbers(s)
        add_page_numbers(s)


def emphasise(par, text: str) -> None:
    """Render **bold**, *italic*, and [n] citation markers as superscripts."""
    for tok in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|\[\d+(?:,\d+)*\])", text):
        if not tok:
            continue
        if re.fullmatch(r"\[\d+(?:,\d+)*\]", tok):
            r = par.add_run(tok.strip("[]"))
            r.font.superscript = True
            continue
        r = par.add_run(re.sub(r"^\*\*|\*\*$|^\*|\*$", "", tok))
        r.bold = tok.startswith("**")
        r.italic = tok.startswith("*") and not tok.startswith("**")


def add_table(doc: Document, block: list[str]) -> None:
    rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in block if not set(r) <= set("|-: ")]
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for r in t.rows:
        r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            para = t.cell(i, j).paragraphs[0]
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = para.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", cell))
            run.font.size = Pt(9)
            run.bold = i == 0
    doc.add_paragraph()


def render(doc: Document, lines: list[str], page_break_before: set[str]) -> None:
    buf: list[str] = []

    def flush():
        if buf:
            emphasise(doc.add_paragraph(), " ".join(buf).strip())
            buf.clear()

    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            add_table(doc, block)
            continue
        if ln.startswith("#"):
            flush()
            title = ln.lstrip("# ").strip()
            if title in page_break_before:
                doc.add_page_break()
            h = doc.add_heading(title, level=min(len(ln) - len(ln.lstrip("#")), 3))
            for r in h.runs:
                r.font.color.rgb = INK
        elif ln.strip() == "---" or not ln.strip():
            flush()
        elif ln.startswith("- "):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            emphasise(p, ln[2:])
        elif re.match(r"^\d+\. ", ln):
            flush()
            p = doc.add_paragraph()
            emphasise(p, ln)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
        else:
            buf.append(ln.strip())
        i += 1
    flush()


def build(with_figures: bool, out: Path) -> None:
    doc = Document()
    style(doc)
    src = SRC.read_text()
    # NDT prescribes References, Tables, Figure Legends, Figures, so the legends are held back
    # and emitted after the tables rather than in the order they appear in the source file.
    cut = src.index("## Figure Legends")
    render(doc, src[:cut].split("\n"), set())

    doc.add_page_break()
    doc.add_heading("Tables", level=2)
    render(doc, TABLES.read_text().split("\n"), set())

    doc.add_page_break()
    render(doc, src[cut:].split("\n"), set())

    if with_figures:
        for n, f in enumerate(FIGURE_FILES, 1):
            if (FIG / f).exists():
                doc.add_page_break()
                doc.add_picture(str(FIG / f), width=Inches(6.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                r = cap.add_run(f"Figure {n}")
                r.italic, r.font.size = True, Pt(10)
    doc.save(out)
    print("wrote", out.name)


def check(src: str) -> None:
    """Fail loudly if the manuscript breaks a stated NDT limit."""
    ab = src[src.index("## Abstract") + 11: src.index("\n---", src.index("## Abstract"))]
    body = src[src.index("## Introduction"): src.index("## Acknowledgements")]
    n_ab, n_body = len(ab.split()), len(body.split())
    refs = len(re.findall(r"^\d+\. ", (PUB / "references.md").read_text(), re.M))
    kw = len(src.split("Keywords: ")[1].split("\n")[0].split(";"))
    head = src.split("Running head: ")[1].split("\n")[0]
    fails = []
    if n_ab > 300:
        fails.append(f"abstract {n_ab} > 300")
    if n_ab + n_body > 3500:
        fails.append(f"abstract+main text {n_ab + n_body} > 3500")
    if refs > 50:
        fails.append(f"references {refs} > 50")
    if kw > 5:
        fails.append(f"keywords {kw} > 5")
    if len(head) > 50:
        fails.append(f"running head {len(head)} characters > 50")
    klp = src[src.index("## Key learning points"): src.index("## Introduction")]
    for name in ("What was known", "This study adds", "Potential impact"):
        if f"**{name}**" not in klp:
            fails.append(f"missing key learning point section: {name}")
    for b in re.findall(r"^- (.+(?:\n  .+)*)", klp, re.M):
        if len(b.split()) > 50:
            fails.append(f"key learning bullet {len(b.split())} words > 50")
    for f in ("Background and hypothesis", "Methods", "Results", "Conclusions"):
        if f"**{f}.**" not in ab:
            fails.append(f"abstract missing required heading: {f}")
    print(f"abstract {n_ab}/300 | abstract+main text {n_ab + n_body}/3500 | references {refs}/50 | "
          f"keywords {kw}/5 | running head {len(head)}/50 chars")
    if fails:
        for f in fails:
            print("FAIL", f)
        raise SystemExit(1)
    print("NDT limits: all OK")


def main() -> None:
    check(SRC.read_text())
    build(False, PUB / "NDT_submission_manuscript.docx")
    build(True, PUB / "NDT_review_copy.docx")
    if SUPP.exists():
        doc = Document()
        style(doc)
        render(doc, SUPP.read_text().split("\n"), set())
        doc.save(PUB / "NDT_supplementary_methods.docx")
        print("wrote NDT_supplementary_methods.docx")


if __name__ == "__main__":
    sys.exit(main())
