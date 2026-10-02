#!/usr/bin/env python3
"""
20_build_ejhg_docx.py

Builds the EJHG submission manuscript file.

Formatting set from the live Guide to Authors (nature.com/ejhg/authors-and-referees/gta,
read 20 September 2026): "Text must be double spaced with a wide margin", unstructured
abstract up to 250 words, main text up to 4,000 words, at most 4 tables and 4 figures,
at most 50 references, and the prescribed section order. Continuous line numbering and
page numbering are added: they are not in the Guide, but editorial offices ask for them
and their absence has caused a return before.

Two files are written:
  EJHG_submission_manuscript.docx   text and tables, no embedded figures (per EJHG,
                                    figures are supplied as separate files)
  EJHG_review_copy.docx             the same with figures embedded, for reading

Author: Christopher Lawrence
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
PUB = ROOT / "publication"
FIG = PUB / "figures_ejhg"
SRC = PUB / "EJHG_manuscript_v2.1.referenced.md"
TABLES = PUB / "tables_ejhg.md"
INK = RGBColor(0, 0, 0)
FIGURE_FILES = ["figure1_discrimination.png", "figure2_avi_attribution.png", "figure3_locus_rank.png"]


def add_line_numbers(section) -> None:
    """Continuous line numbering, restarting at each page is not used."""
    sectPr = section._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    sectPr.append(ln)


def add_page_numbers(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text, kind in (("PAGE", "begin"), (None, None)):
        pass
    run = p.add_run()
    for instr, typ in (("begin", None), (None, "PAGE"), ("end", None)):
        el = OxmlElement("w:fldChar") if instr else OxmlElement("w:instrText")
        if instr:
            el.set(qn("w:fldCharType"), instr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        run._r.append(el)


def style(doc: Document) -> None:
    n = doc.styles["Normal"]
    n.font.name, n.font.size, n.font.color.rgb = "Times New Roman", Pt(12), INK
    pf = n.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(0)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.25)
        s.top_margin = s.bottom_margin = Inches(1.0)
        add_line_numbers(s)
        add_page_numbers(s)


def emphasise(par, text: str) -> None:
    for tok in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)", text):
        if not tok:
            continue
        r = par.add_run(re.sub(r"^\*\*|\*\*$|^\*|\*$|^`|`$", "", tok))
        r.bold = tok.startswith("**")
        r.italic = tok.startswith("*") and not tok.startswith("**")


def add_table(doc: Document, block: list[str]) -> None:
    rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in block if not set(r) <= set("|-: ")]
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))  # repeat header across pages
    for r in t.rows:
        # w:cantSplit stops a single row breaking across a page, which had split Table 4's long
        # assay descriptions over two pages in the previous build.
        cs = OxmlElement("w:cantSplit")
        r._tr.get_or_add_trPr().append(cs)
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            para = t.cell(i, j).paragraphs[0]
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = para.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", cell))
            run.font.size = Pt(9)
            run.bold = i == 0
    doc.add_paragraph()


def build(with_figures: bool, out: Path) -> None:
    doc = Document()
    style(doc)
    lines = SRC.read_text().split("\n")
    buf: list[str] = []

    def flush():
        if buf:
            emphasise(doc.add_paragraph(), " ".join(buf).strip())
            buf.clear()

    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            add_table(doc, block)
            continue
        if ln.startswith("#"):
            flush()
            if ln.lstrip("# ").strip() == "Figure Legends":
                doc.add_page_break()
            h = doc.add_heading(ln.lstrip("# ").strip(), level=min(len(ln) - len(ln.lstrip("#")), 3))
            for r in h.runs:
                r.font.color.rgb = INK
        elif ln.strip() == "---":
            flush()
        elif re.match(r"^\d+\. ", ln):
            flush()
            p = doc.add_paragraph()
            emphasise(p, ln)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
        elif not ln.strip():
            flush()
        else:
            buf.append(ln.strip())
        i += 1
    flush()

    # tables, after the text, as EJHG requests
    doc.add_page_break()
    doc.add_heading("Tables", level=2)
    tl = TABLES.read_text().split("\n")
    i = 0
    while i < len(tl):
        if tl[i].startswith("|"):
            block = []
            while i < len(tl) and tl[i].startswith("|"):
                block.append(tl[i]); i += 1
            add_table(doc, block)
            continue
        if tl[i].strip():
            emphasise(doc.add_paragraph(), tl[i].strip())
        i += 1

    if with_figures:
        doc.add_page_break()
        doc.add_heading("Figures", level=2)
        for n, f in enumerate(FIGURE_FILES, 1):
            if (FIG / f).exists():
                if n > 1:
                    doc.add_page_break()
                doc.add_picture(str(FIG / f), width=Inches(6.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                r = cap.add_run(f"Figure {n}")
                r.italic, r.font.size = True, Pt(10)
    doc.save(out)
    print("wrote", out.name)


def main() -> None:
    build(False, PUB / "EJHG_submission_manuscript.docx")
    build(True, PUB / "EJHG_review_copy.docx")


if __name__ == "__main__":
    main()
