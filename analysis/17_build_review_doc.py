#!/usr/bin/env python3
"""
17_build_review_doc.py

Builds a single Word document for review: the manuscript text, its tables as real
Word tables, and the figures embedded in place with captions. This is a reading
copy, not the submission file; the submission file is assembled to the journal's
own format.

Usage: python3 17_build_review_doc.py

Output: publication/PKD_AlphaGenome_REVIEW_COPY.docx

Author: Christopher Lawrence
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
PUB = ROOT / "publication"
FIG = PUB / "figures"
SRC = PUB / "DRAFT_manuscript_v1.0.referenced.md"
OUT = PUB / "PKD_AlphaGenome_REVIEW_COPY.docx"

# Where each figure goes, and its caption.
FIGURES = {
    "(Figure 1, Table 2.)": ("fig1_scores_by_label.png",
                    "Figure 1. What each tool scores variants that were tested in the laboratory. "
                    "Each point is one variant; the black bar is the median. AVI is a composite "
                    "pathogenicity score, so its panel is restricted to non-coding and synonymous "
                    "variants, where a high score cannot arise from a protein consequence."),
    "(Figure 2)": ("fig2_blind_spots.png",
                    "Figure 2. Five variants that act after the message is transcribed. Bars show "
                    "AlphaGenome's predicted change in PKD1 RNA; a doubling would be 1.0 on this "
                    "axis. The measured effect is given as text because two of the five sources "
                    "report a direction only."),
    "(Figure 3)": ("fig3_avi_attribution.png",
                   "Figure 3. What drives the AVI score, by variant class. Conservation carries the "
                   "two drug-target sites; protein damage carries the coding positives."),
    "(Figure 4.)": ("fig4_locus_rank.png",
                    "Figure 4. How far down the gene-wide list each variant sits: the percentage of "
                    "all possible PKD1 variants scoring higher, for each tool, on a log scale. "
                    "Points below the diagonal are ranked higher by SpliceAI than by AVI."),
}
INK = RGBColor(0x0B, 0x0B, 0x0B)


def add_md_table(doc: Document, block: list[str]) -> None:
    rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in block if not set(r) <= set("|-: ")]
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            para = tbl.cell(i, j).paragraphs[0]
            run = para.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", cell))
            run.font.size = Pt(9)
            if i == 0:
                run.bold = True
    doc.add_paragraph()


def emphasise(par, text: str) -> None:
    """Render **bold**, *italic* and `code` inside a paragraph."""
    for tok in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)", text):
        if not tok:
            continue
        run = par.add_run(re.sub(r"^\*\*|\*\*$|^\*|\*$|^`|`$", "", tok))
        run.bold = tok.startswith("**")
        run.italic = tok.startswith("*") and not tok.startswith("**")
        if tok.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(9)


def main() -> None:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size, st.font.color.rgb = "Calibri", Pt(11), INK

    lines = SRC.read_text().split("\n")
    i, buf = 0, []

    def flush():
        if buf:
            p = doc.add_paragraph()
            emphasise(p, " ".join(buf).strip())
            buf.clear()

    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            add_md_table(doc, block)
            continue
        if ln.startswith("#"):
            flush()
            lvl = len(ln) - len(ln.lstrip("#"))
            h = doc.add_heading(ln.lstrip("# ").strip(), level=min(lvl, 4))
            for r in h.runs:
                r.font.color.rgb = INK
        elif ln.strip() in ("---",):
            flush()
        elif ln.startswith(("- ", "* ")):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            emphasise(p, ln[2:])
        elif re.match(r"^\d+\. ", ln):
            flush()
            p = doc.add_paragraph(style="List Number")
            emphasise(p, re.sub(r"^\d+\. ", "", ln))
        elif not ln.strip():
            flush()
        else:
            buf.append(ln.strip())
        # figure insertion points
        for marker, (png, cap) in FIGURES.items():
            if marker in ln and (FIG / png).exists():
                flush()
                doc.add_picture(str(FIG / png), width=Inches(6.3))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                c = doc.add_paragraph()
                r = c.add_run(cap)
                r.font.size, r.italic = Pt(9), True
        i += 1
    flush()

    doc.add_page_break()
    doc.add_heading("Supplementary Table S1. Sources of the truth set", level=2)
    for ln in (PUB / "supp_table_S1_sources.md").read_text().split("\n"):
        if ln.startswith("#") or not ln.strip():
            continue
        p = doc.add_paragraph()
        emphasise(p, ln)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.size = Pt(9)
    doc.save(OUT)
    print("wrote", OUT.name)


if __name__ == "__main__":
    main()
